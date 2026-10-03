# -*- coding: utf-8 -*-
"""Repair API Design workbook Endpoint column: unique paths per row (EN + VI sheets)."""
from pathlib import Path

import openpyxl

OUT_PATH = Path(r"C:\Users\PC\Downloads\API Design - OOP AutoGrader.xlsx")

# Match (router, HTTP method, English description prefix) -> endpoint path (column B)
# Descriptions are from the English sheet; Vietnamese sheet is updated in parallel by row.
ENDPOINT_BY_DESC = [
    ("/api/users", "GET", "Get one user", "/getUser/{id}"),
    ("/api/users", "DELETE", "Hard-delete a user", "/deleteUser/{id}"),
    ("/api/users", "PUT", "Update a user's profile", "/updateUser/{id}"),
    ("/api/labs", "GET", "List labs in the current term", "/list"),
    ("/api/labs", "GET", "Get submission-count", "/{labId}/stats?studentId"),
    ("/api/labs", "GET", "Lecturer lab analytics", "/{labId}/statistics"),
    ("/api/labs", "GET", "Paginated roster of students who submitted", "/{labId}/submissions?page&size&sort&afterName&afterId&search"),
    ("/api/labs", "GET", "Full submitter roster for a lab", "/{labId}/submissions/export?sort"),
    ("/api/labs", "GET", "Lab attempt history for one student", "/{labId}/students/{studentId}/attempts"),
    ("/api/labs/{labId}/challenges", "GET", "List challenges (names", "/list"),
    ("/api/labs/{labId}/challenges", "GET", "Get MMD (class diagram)", "/{challengeId}/mmd?studentId&submissionId"),
    ("/api/labs/{labId}/challenges", "GET", "Get Class tab grading", "/{challengeId}/class?studentId&submissionId"),
    ("/api/labs/{labId}/challenges", "GET", "Get operational testcase", "/{challengeId}/testcases?studentId&submissionId"),
    ("/api/labs/{labId}/challenges", "GET", "Per-challenge submission stats", "/{challengeId}/stats?studentId"),
    ("/api/labs/{labId}/challenges", "GET", "Paginated roster of students with a graded submission", "/{challengeId}/students?page&size&sort"),
    ("/api/lecturer/labs", "POST", "Deep-clone selected labs", "/clone"),
    ("/api/lecturer/labs", "POST", "Create a new, empty lab", "/create"),
    ("/api/lecturer/labs", "POST", "Dry-run one testcase", "/{labId}/challenges/{challengeId}/testcases/dry-run"),
    ("/api/lecturer/terms", "GET", "List all academic terms", "/list"),
    ("/api/lecturer/terms", "POST", "Create a new term", "/create"),
    ("/api/lecturer/terms", "POST", "Mark a term as the current", "/{termId}/current"),
    ("/api/lecturer/terms", "GET", "Get enrolled and available students", "/{termId}/roster"),
    ("/api/lecturer/terms", "GET", "List students enrolled in a term", "/{termId}/students"),
    ("/api/lecturer/terms", "GET", "List students not yet enrolled", "/{termId}/available-students"),
    ("/api/lecturer/terms", "POST", "Enroll a list of students", "/{termId}/students"),
    ("/api/lecturer/terms", "POST", "Bulk-enroll students", "/{termId}/students/import"),
]


def resolve_endpoint(router: str, method: str, desc: str) -> str | None:
    if not desc:
        return None
    desc = str(desc).strip()
    for r, m, prefix, endpoint in ENDPOINT_BY_DESC:
        if router == r and method == m and desc.startswith(prefix):
            return endpoint
    return None


def patch_workbook(path: Path) -> None:
    wb = openpyxl.load_workbook(path)
    for sheet_name in ("English", "Vietnamese"):
        ws = wb[sheet_name]
        router = None
        desc_col = 3  # English description; Vietnamese uses same column for matching via EN sheet
        en_ws = wb["English"] if sheet_name == "Vietnamese" else ws

        for row in range(2, ws.max_row + 1):
            router_cell = ws.cell(row=row, column=1).value
            if router_cell:
                router = str(router_cell).strip()

            method = ws.cell(row=row, column=4).value
            if not method or not router:
                continue
            method = str(method).strip()

            match_desc = en_ws.cell(row=row, column=desc_col).value
            endpoint = resolve_endpoint(router, method, match_desc)
            if endpoint:
                ws.cell(row=row, column=2).value = endpoint

            # User path renames in notes / other columns
            for col in range(1, 12):
                val = ws.cell(row=row, column=col).value
                if isinstance(val, str):
                    val = (
                        val.replace("PUT /api/users/{id}", "PUT /api/users/updateUser/{id}")
                        .replace("DELETE /api/users/{id}", "DELETE /api/users/deleteUser/{id}")
                        .replace("GET /api/users/{id}", "GET /api/users/getUser/{id}")
                    )
                    ws.cell(row=row, column=col).value = val

    wb.save(path)
    print(f"Patched endpoints: {path}")


def main() -> None:
    if not OUT_PATH.exists():
        raise SystemExit(f"Workbook not found: {OUT_PATH}")
    patch_workbook(OUT_PATH)


if __name__ == "__main__":
    main()
